"""Unisce i dati di tutte le sessioni e produce le tabelle dell'analisi.

    python scripts/analisi.py                    # tutte le sessioni
    python scripts/analisi.py --sessioni P01 P02

Legge, per ogni sessione:
  dati/dati_<C>.xlsx   i questionari inseriti a mano (4 fogli)
  logs/*.csv           quello che ha scritto il sistema durante la partita

e scrive in analisi/:
  osservazioni.csv     una riga per partecipante x condizione  <- la tabella principale
  partecipanti.csv     una riga per partecipante (demografia + questionario finale)
  item.csv             formato lungo: partecipante x condizione x item
  aperte.csv           le risposte a testo libero
  riepilogo.txt        il riassunto stampato a video

Nessun test statistico: con pochi partecipanti non avrebbe senso. Qui si costruisce
la base dati e si controlla che regga.
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

RADICE = Path(__file__).resolve().parent.parent
SESSIONI = RADICE / "sessioni"
USCITA = RADICE / "analisi"

# --- corrispondenze fissate in esperimento/chiave_spoglio.md --------------
MERCANTE_CONDIZIONE = {
    "adrianne_avenicci": "V",
    "arcadia": "N",
    "belethor": "A",
    "lucan_valerius": "D",
}
ORDINE_CONDIZIONI = ["V", "N", "A", "D"]

ITEM_INVERSI = {"i02"}          # punteggio corretto = 8 - risposta

DIMENSIONI = {
    "comportamento": ["i01", "i02"],
    "relazione_sociale": ["i03"],
    "intelligenza": ["i04", "i05", "i06"],
    "emozione": ["i07", "i08"],
    "memoria_percepita": ["i09", "i10", "i11"],
    "direzione_relazione": ["i12"],
    "collera_percepita": ["i13"],
    "trattativa_percepita": ["i14"],
    "credibilita": ["i15", "i16", "i17"],
}

# b13 e b14 restano fuori: sono il costo dichiarato, non si sommano agli altri
BLOCCHI_B = {
    "valore_soggettivo": ["b01", "b02", "b03", "b04"],
    "coinvolgimento": ["b05", "b06", "b07", "b08"],
    "preferenza_meccanica": ["b09", "b10", "b11", "b12"],
}


def numero(v):
    """Converte in float. Passando da Google Sheets i punteggi tornano come testo."""
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    testo = str(v).strip().replace(",", ".")
    try:
        return float(testo)
    except ValueError:
        return None


def media(valori):
    valori = [v for v in valori if v is not None]
    return round(sum(valori) / len(valori), 3) if valori else None


# ---------------------------------------------------------------- Excel ---

def _campi_per_chiave(ws, righe):
    """Coppie (colonna D = nome tecnico, colonna B = valore inserito)."""
    fuori = {}
    for r in righe:
        chiave = ws.cell(row=r, column=4).value
        if chiave:
            fuori[str(chiave)] = ws.cell(row=r, column=2).value
    return fuori


def leggi_foglio(percorso: Path):
    wb = load_workbook(percorso, data_only=True)

    sessione = _campi_per_chiave(wb["1 Sessione"], range(5, 40))
    finale = _campi_per_chiave(wb["3 Finale"], range(5, 60))

    # la build sta nel titolo: "Sessione P02 — build 2 — scheda partecipante"
    build = None
    titolo = wb["1 Sessione"]["A1"].value or ""
    trovato = re.search(r"build\s+(\d)", str(titolo))
    if trovato:
        build = int(trovato.group(1))

    ws = wb["2 Questionari"]
    # le due righe grigie in fondo dicono, colonna per colonna, condizione e mercante
    colonne: dict[int, dict] = {}
    for r in range(5, ws.max_row + 1):
        etichetta = str(ws.cell(row=r, column=2).value or "")
        if etichetta.startswith("condizione"):
            for c in range(3, 7):
                colonne.setdefault(c, {})["condizione"] = ws.cell(row=r, column=c).value
        elif etichetta.startswith("mercante"):
            for c in range(3, 7):
                colonne.setdefault(c, {})["mercante"] = ws.cell(row=r, column=c).value

    item = []
    for r in range(5, 22):
        chiave = ws.cell(row=r, column=1).value
        if not (isinstance(chiave, str) and chiave.startswith("i")):
            continue
        for c in range(3, 7):
            meta = colonne.get(c, {})
            if not meta.get("condizione"):
                continue
            item.append({
                "condizione": meta["condizione"],
                "mercante": meta["mercante"],
                "posizione": c - 2,
                "item": chiave,
                "valore": numero(ws.cell(row=r, column=c).value),
            })

    ws = wb["4 Aperte"]
    aperte = []
    for r in range(5, ws.max_row + 1):
        domanda = ws.cell(row=r, column=1).value
        risposta = ws.cell(row=r, column=2).value
        if domanda and risposta:
            aperte.append({"domanda": str(domanda), "risposta": str(risposta).strip()})

    return sessione, finale, item, aperte, build


# ------------------------------------------------------------------ log ---

def leggi_log(cartella: Path, nome: str) -> pd.DataFrame:
    percorso = cartella / "logs" / nome
    if not percorso.exists():
        return pd.DataFrame()
    # parser CSV vero: le battute degli NPC contengono virgole e a capo
    return pd.read_csv(percorso, encoding="utf-8")


def esiti_dai_turni(cartella: Path) -> dict:
    """Esito finale per mercante, letto dai turni di trattativa.

    Serve per gli abbandoni: una trattativa chiusa senza accordo non produce
    alcuna riga fra le transazioni, quindi senza questo il tasso di abbandono
    resterebbe invisibile e la condizione apparirebbe come dato mancante.
    """
    turni = leggi_log(cartella, "negotiation_turns_log.csv")
    if turni.empty or "status" not in turni.columns:
        return {}
    fuori = {}
    for mercante, gruppo in turni.groupby("npc_id"):
        chiusure = gruppo[gruppo["status"].isin(["deal", "walkaway"])]
        if chiusure.empty:
            continue
        fuori[mercante] = {
            "esito": chiusure.iloc[-1]["status"],
            "round": int(gruppo["turn"].max()) if "turn" in gruppo else None,
        }
    return fuori


def dati_comportamentali(cartella: Path) -> dict:
    esiti = esiti_dai_turni(cartella)
    tx = leggi_log(cartella, "merchant_transactions_log.csv")
    if tx.empty:
        # nessun affare concluso: restano gli eventuali abbandoni
        return {MERCANTE_CONDIZIONE[m]: {"condizione": MERCANTE_CONDIZIONE[m],
                                         "mercante": m, "esito": d["esito"],
                                         "round": d["round"], "prezzo_pattuito": None,
                                         "riserva": None, "scarto_riserva": None}
                for m, d in esiti.items() if m in MERCANTE_CONDIZIONE}
    tx = tx.rename(columns={
        "npc_id": "mercante",
        "npc_target": "riserva",
        "agreed_price": "prezzo_pattuito",
        "turns": "round",
        "outcome": "esito",
    })
    tx["condizione"] = tx["mercante"].map(MERCANTE_CONDIZIONE)
    tx["scarto_riserva"] = tx["prezzo_pattuito"] - tx["riserva"]

    mod = leggi_log(cartella, "price_modifier_evolution.csv")
    if not mod.empty:
        mod = mod.rename(columns={"npc_id": "mercante"})
        colonne = [c for c in ("mercante", "price_modifier", "relationship_stage",
                               "affinity", "respect", "suspicion", "gratitude")
                   if c in mod.columns]
        tx = tx.merge(mod[colonne].drop_duplicates("mercante"), on="mercante", how="left")

    fuori = {r["condizione"]: dict(r) for _, r in tx.iterrows()}

    # I mercanti che compaiono nei turni ma non fra le transazioni sono abbandoni.
    for mercante, d in esiti.items():
        cond = MERCANTE_CONDIZIONE.get(mercante)
        if cond and cond not in fuori:
            fuori[cond] = {"condizione": cond, "mercante": mercante,
                           "esito": d["esito"], "round": d["round"],
                           "prezzo_pattuito": None, "riserva": None,
                           "scarto_riserva": None}
    return fuori


# ---------------------------------------------------------- assemblaggio ---

def raccogli(codici=None):
    cartelle = sorted(d for d in SESSIONI.glob("P*") if d.is_dir())
    if codici:
        cartelle = [d for d in cartelle if d.name in codici]
    if not cartelle:
        sys.exit("nessuna sessione trovata in sessioni/")

    righe_oss, righe_part, righe_item, righe_aperte = [], [], [], []

    for cartella in cartelle:
        codice = cartella.name
        fogli = list((cartella / "dati").glob("*.xlsx"))
        if not fogli:
            print(f"  {codice}: nessun foglio dati, salto")
            continue

        sessione, finale, item, aperte, build = leggi_foglio(fogli[0])

        # --- una riga per partecipante
        p = {"partecipante": codice, "build": build}
        p.update(sessione)
        for chiave, valore in finale.items():
            p[chiave] = numero(valore) if chiave.startswith("b") else valore
        for nome, chiavi in BLOCCHI_B.items():
            p[nome] = media([numero(finale.get(k)) for k in chiavi])
        p["attrito_latenza"] = numero(finale.get("b13"))
        p["attrito_riconoscimento"] = numero(finale.get("b14"))
        righe_part.append(p)

        # --- item in formato lungo
        for riga in item:
            righe_item.append({"partecipante": codice, **riga})

        # --- punteggi, con l'item inverso ricodificato
        punteggi: dict[str, dict[str, float]] = {}
        for riga in item:
            valore = riga["valore"]
            if valore is None:
                continue
            if riga["item"] in ITEM_INVERSI:
                valore = 8 - valore
            punteggi.setdefault(riga["condizione"], {})[riga["item"]] = valore

        comportamentali = dati_comportamentali(cartella)
        prezzo_vanilla = numero(sessione.get("prezzo_vanilla"))
        meta_condizione = {r["condizione"]: r for r in item}

        for cond in ORDINE_CONDIZIONI:
            if cond not in meta_condizione:
                continue
            meta = meta_condizione[cond]
            oss = {
                "partecipante": codice,
                "build": build,
                "condizione": cond,
                "mercante": meta["mercante"],
                "posizione": meta["posizione"],
            }
            for nome, chiavi in DIMENSIONI.items():
                oss[nome] = media([punteggi.get(cond, {}).get(k) for k in chiavi])

            comp = comportamentali.get(cond)
            if comp is not None:
                for campo in ("riserva", "prezzo_pattuito", "scarto_riserva", "round",
                              "esito", "price_modifier", "relationship_stage"):
                    if campo in comp:
                        oss[campo] = comp[campo]
            elif cond == "V":
                # il baratto vanilla non lascia log: il prezzo e' quello annotato a mano
                oss["prezzo_pattuito"] = prezzo_vanilla
                oss["riserva"] = None
                oss["scarto_riserva"] = None
                oss["round"] = 0
                oss["esito"] = "deal"
            righe_oss.append(oss)

        for a in aperte:
            righe_aperte.append({"partecipante": codice, **a})

    return (pd.DataFrame(righe_oss), pd.DataFrame(righe_part),
            pd.DataFrame(righe_item), pd.DataFrame(righe_aperte))


# ------------------------------------------------------------- riepilogo ---

def _riga_controllo(nome, esito, valori):
    stato = "OK        " if esito else "NON regge "
    return f"  {nome:34} {stato} " + "  ".join(valori)


def riepilogo(oss: pd.DataFrame, part: pd.DataFrame) -> str:
    r = []
    r.append(f"PARTECIPANTI: {part['partecipante'].nunique()}    "
             f"osservazioni: {len(oss)}")
    if "build" in part:
        r.append(f"Build coperte: {sorted(int(b) for b in part['build'].dropna().unique())}")
    if "esperienza_skyrim" in part:
        r.append(f"Esperienza con Skyrim: {part['esperienza_skyrim'].value_counts().to_dict()}")
    r.append("")

    con_questionario = sorted(oss.loc[oss["memoria_percepita"].notna(), "partecipante"].unique())
    senza = [p for p in sorted(oss["partecipante"].unique()) if p not in con_questionario]
    if senza:
        r.append(f"ATTENZIONE - questionari ancora da inserire: {', '.join(senza)}")
        r.append("  (le medie di percezione sono calcolate solo su chi e' gia' spogliato)")
        r.append("")

    r.append(f"--- COMPORTAMENTO ECONOMICO (media per condizione, n={len(oss['partecipante'].unique())})")
    colonne = [c for c in ("prezzo_pattuito", "riserva", "scarto_riserva", "round") if c in oss]
    if colonne:
        tabella = oss.groupby("condizione")[colonne].mean().reindex(ORDINE_CONDIZIONI)
        r.append(tabella.round(2).to_string())
    r.append("")

    if "esito" in oss:
        abb = oss.assign(abbandono=(oss["esito"] == "walkaway"))
        conteggio = abb.groupby("condizione")["abbandono"].sum().reindex(ORDINE_CONDIZIONI)
        totali = abb.groupby("condizione")["esito"].count().reindex(ORDINE_CONDIZIONI)
        r.append("--- ABBANDONI (trattative chiuse senza accordo)")
        r.append("  " + "   ".join(
            f"{c}: {int(conteggio.get(c, 0))}/{int(totali.get(c, 0))}"
            for c in ORDINE_CONDIZIONI if not pd.isna(totali.get(c))))
        r.append("")

    r.append(f"--- PERCEZIONE (media per condizione, scala 1-7, n={len(con_questionario)})")
    dim = [d for d in DIMENSIONI if d in oss]
    tabella = oss.groupby("condizione")[dim].mean().reindex(ORDINE_CONDIZIONI)
    r.append(tabella.round(2).to_string())
    r.append("")

    r.append("--- CONTROLLI PREVISTI DAL PIANO")
    mem = oss.groupby("condizione")["memoria_percepita"].mean()
    esito = (mem.get("A", 0) > mem.get("N", 9) and mem.get("D", 0) > mem.get("N", 9)
             and mem.get("N", 0) > mem.get("V", 9))
    r.append(_riga_controllo("memoria percepita  A,D > N > V", esito,
                             [f"{c}={mem.get(c, float('nan')):.2f}" for c in ORDINE_CONDIZIONI]))

    tr = oss.groupby("condizione")["trattativa_percepita"].mean()
    altri = [tr.get(c, 0) for c in ("N", "A", "D")]
    esito = tr.get("V", 9) < min(altri) if altri else False
    r.append(_riga_controllo("controllo manipolazione: V < resto", esito,
                             [f"{c}={tr.get(c, float('nan')):.2f}" for c in ORDINE_CONDIZIONI]))

    col = oss.groupby("condizione")["collera_percepita"].mean()
    esito = col.get("D", 0) > max(col.get(c, 0) for c in ("V", "N", "A"))
    r.append(_riga_controllo("collera percepita: D la piu' alta", esito,
                             [f"{c}={col.get(c, float('nan')):.2f}" for c in ORDINE_CONDIZIONI]))

    if "prezzo_pattuito" in oss:
        pr = oss.groupby("condizione")["prezzo_pattuito"].mean()
        esito = pr.get("A", 9e9) < pr.get("N", 0) < pr.get("D", 0)
        r.append(_riga_controllo("prezzo pattuito  A < N < D", esito,
                                 [f"{c}={pr.get(c, float('nan')):.2f}" for c in ("A", "N", "D")]))
    r.append("")

    r.append("--- QUESTIONARIO FINALE (conteggi)")
    etichette = [("a1_differenze", "ha notato differenze"),
                 ("a4_ostile", "il piu' ostile"),
                 ("a4_amichevole", "il piu' amichevole"),
                 ("a5_ricordava", "sembrava ricordare"),
                 ("a6_coinvolgente", "trattativa piu' coinvolgente"),
                 ("a7_attribuzione", "attribuzione del prezzo")]
    for campo, testo in etichette:
        if campo in part:
            r.append(f"  {testo:32} {part[campo].value_counts().to_dict()}")
    r.append("")

    r.append("--- VALUTAZIONE DELLA MECCANICA (media, 1-7)")
    for nome in list(BLOCCHI_B) + ["attrito_latenza", "attrito_riconoscimento"]:
        if nome in part:
            valori = part[nome].dropna()
            if len(valori):
                r.append(f"  {nome:26} {valori.mean():.2f}")
    return "\n".join(r)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--sessioni", nargs="*", help="codici da includere (default: tutti)")
    args = ap.parse_args()

    oss, part, item, aperte = raccogli(args.sessioni)
    USCITA.mkdir(exist_ok=True)
    oss.to_csv(USCITA / "osservazioni.csv", index=False, encoding="utf-8")
    part.to_csv(USCITA / "partecipanti.csv", index=False, encoding="utf-8")
    item.to_csv(USCITA / "item.csv", index=False, encoding="utf-8")
    aperte.to_csv(USCITA / "aperte.csv", index=False, encoding="utf-8")

    testo = riepilogo(oss, part)
    (USCITA / "riepilogo.txt").write_text(testo, encoding="utf-8")
    print(testo)
    print()
    print("scritti in analisi/: osservazioni.csv, partecipanti.csv, item.csv, "
          "aperte.csv, riepilogo.txt")


if __name__ == "__main__":
    main()
