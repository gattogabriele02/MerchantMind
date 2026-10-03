"""Crea il foglio Excel in cui inserire a mano i dati di UNA sessione.

    python scripts/crea_modulo_dati.py P02 2

Produce  sessioni/P02/dati/dati_P02.xlsx  con quattro fogli:
  1 Sessione       demografia, prezzo vanilla, durata, note
  2 Questionari    17 item x 4 condizioni, gia' etichettate secondo la build
  3 Finale         questionario finale, parti A e B
  4 Aperte         risposte a testo libero

Il testo di ogni domanda e' accanto alla casella: si compila leggendo il foglio
di carta e digitando il numero nella riga corrispondente, senza dover ricordare
che cosa fosse l'item 11.

Perche' Excel e non CSV: Excel in italiano usa il punto e virgola come
separatore di lista, quindi un CSV separato da virgole finisce tutto in una
cella sola. Il convertitore (converti_dati.py) rilegge questi fogli e produce i
CSV per l'analisi, cosi' il formato di scrittura resta comodo e quello di
lettura resta pulito.
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

RADICE = Path(__file__).resolve().parent.parent

SEQUENZE = {
    1: ("V", "N", "A", "D"),
    2: ("N", "D", "V", "A"),
    3: ("A", "V", "D", "N"),
    4: ("D", "A", "N", "V"),
}
MERCANTI = {
    "V": "adrianne_avenicci",
    "N": "arcadia",
    "A": "belethor",
    "D": "lucan_valerius",
}
NOMI = {
    "V": "Adrianne Avenicci",
    "N": "Arcadia",
    "A": "Belethor",
    "D": "Lucan Valerius",
}

# I 17 item del questionario per mercante, nell'ordine del modulo stampato.
ITEM = [
    ("i01", "Si e' comportato in modo coerente e naturale."),
    ("i02", "A volte si e' comportato in modo fuori luogo."),
    ("i03", "Ha interagito con me sul piano sociale, non solo commerciale."),
    ("i04", "Sembrava capace di farsi una strategia."),
    ("i05", "Sembrava aver imparato da esperienze passate."),
    ("i06", "Sembrava avere memoria."),
    ("i07", "Sembrava capace di provare emozioni."),
    ("i08", "Le emozioni che ha mostrato erano adatte alla situazione."),
    ("i09", "Sembrava ricordare cose della nostra storia passata."),
    ("i10", "Ha citato fatti o esperienze che avevamo in comune."),
    ("i11", "Il suo modo di trattarmi dipendeva da com'erano andate le cose fra noi."),
    ("i12", "Mi ha trattato meglio di come tratterebbe uno sconosciuto."),
    ("i13", "Era irritato o in collera con me."),
    ("i14", "Ho potuto influenzare il prezzo che ho pagato."),
    ("i15", "Era credibile."),
    ("i16", "Si comportava come una persona vera."),
    ("i17", "L'interazione con lui mi e' piaciuta."),
]

# I 14 item della parte B del questionario finale.
ITEM_B = [
    ("b01", "Sono soddisfatto dei prezzi che ho spuntato."),
    ("b02", "Sono soddisfatto di come mi sono comportato mentre trattavo."),
    ("b03", "Mi sono sentito trattato in modo equo."),
    ("b04", "Ho capito che cosa conveniva fare per ottenere un prezzo migliore."),
    ("b05", "Ero preso da quello che stavo facendo."),
    ("b06", "Mentre trattavo mi dimenticavo di quello che avevo intorno."),
    ("b07", "Contrattare a voce mi ha fatto sentire dentro il mondo del gioco."),
    ("b08", "Trattando ho avuto l'impressione di avere un rapporto con il mercante."),
    ("b09", "Quello che dicevo cambiava il modo in cui il mercante si comportava."),
    ("b10", "Comprare parlando e' piu' interessante che scegliere da un menu."),
    ("b11", "Preferirei comprare parlando, se potessi scegliere."),
    ("b12", "Mi piacerebbe trovare mercanti cosi' nei giochi a cui gioco."),
    ("b13", "Le attese fra una battuta e l'altra mi hanno infastidito."),
    ("b14", "Mi e' capitato di dover ripetere perche' non venivo capito."),
]

# Campi della scheda di sessione: (chiave, etichetta, valori ammessi o None)
CAMPI_SESSIONE = [
    ("eta", "Eta (anni)", None),
    ("genere", "Genere", ["uomo", "donna", "non_binario", "non_risponde"]),
    ("lingua_madre", "Lingua madre", None),
    ("ore_gioco", "Ore di gioco a settimana", ["mai", "<2", "2-5", "5-10", ">10"]),
    ("esperienza_skyrim", "Hai mai giocato a Skyrim?", ["mai", "qualche_ora", "molte_ore"]),
    ("conosce_commercio", "Conosce il commercio di Skyrim?", ["no", "sentito", "si"]),
    ("altri_gdr", "Altri giochi di ruolo con mercanti", ["no", "qualcuno", "molti"]),
    ("assistenti_vocali", "Usa assistenti vocali", ["mai", "raramente", "ogni_tanto", "spesso"]),
    ("parlato_npc", "Ha gia' parlato a voce con un NPC", ["mai", "una_due", "piu_volte"]),
    ("prezzo_vanilla", "Prezzo pagato ad Adrianne (condizione V)", None),
    ("durata_minuti", "Durata della sessione (minuti)", None),
    ("note", "Note del conduttore", None),
]

TITOLO = Font(bold=True, size=12)
INTESTAZIONE = Font(bold=True, color="FFFFFF")
SFONDO_INTESTAZIONE = PatternFill("solid", fgColor="44546A")
SFONDO_DA_COMPILARE = PatternFill("solid", fgColor="FFF2CC")
SFONDO_FISSO = PatternFill("solid", fgColor="EDEDED")
A_CAPO = Alignment(wrap_text=True, vertical="top")


def _intestazione(ws, riga, valori, larghezze=None):
    for col, testo in enumerate(valori, start=1):
        c = ws.cell(row=riga, column=col, value=testo)
        c.font = INTESTAZIONE
        c.fill = SFONDO_INTESTAZIONE
        c.alignment = A_CAPO
    if larghezze:
        for col, w in enumerate(larghezze, start=1):
            ws.column_dimensions[get_column_letter(col)].width = w


def _validazione_1_7(ws):
    dv = DataValidation(
        type="whole", operator="between", formula1=1, formula2=7,
        allow_blank=True, showErrorMessage=True,
        errorTitle="Valore non valido", error="La scala va da 1 a 7.",
    )
    ws.add_data_validation(dv)
    return dv


def _validazione_elenco(ws, valori):
    dv = DataValidation(
        type="list", formula1='"' + ",".join(valori) + '"',
        allow_blank=True, showErrorMessage=True,
    )
    ws.add_data_validation(dv)
    return dv


def foglio_sessione(wb, codice, build):
    ws = wb.create_sheet("1 Sessione")
    ws["A1"] = f"Sessione {codice} — build {build} — scheda partecipante"
    ws["A1"].font = TITOLO
    ws["A2"] = "Dalla pagina 1 del modulo cartaceo. Compila la colonna gialla."

    _intestazione(ws, 4, ["Campo", "Valore", "Valori ammessi"], [38, 22, 40])
    riga = 5
    for chiave, etichetta, ammessi in CAMPI_SESSIONE:
        ws.cell(row=riga, column=1, value=etichetta).alignment = A_CAPO
        cella = ws.cell(row=riga, column=2)
        cella.fill = SFONDO_DA_COMPILARE
        if chiave == "prezzo_vanilla":
            cella.value = 50
            ws.cell(row=riga, column=3, value="fisso: 50")
        elif ammessi:
            _validazione_elenco(ws, ammessi).add(cella)
            ws.cell(row=riga, column=3, value=" / ".join(ammessi)).alignment = A_CAPO
        ws.cell(row=riga, column=4, value=chiave).font = Font(color="999999", size=8)
        riga += 1

    ws.freeze_panes = "A5"
    return ws


def foglio_questionari(wb, codice, build):
    ws = wb.create_sheet("2 Questionari")
    ws["A1"] = f"Sessione {codice} — questionario per mercante (pagine 2-5)"
    ws["A1"].font = TITOLO
    ws["A2"] = ("Una colonna per mercante, nell'ordine in cui li ha incontrati. "
                "Scala 1-7. L'item 2 e' inverso: NON ricodificarlo, si fa in analisi.")
    ws["A2"].alignment = A_CAPO

    sequenza = SEQUENZE[build]
    intestazioni = ["Item", "Domanda"]
    for pos, cond in enumerate(sequenza, start=1):
        intestazioni.append(f"{pos}. {NOMI[cond]}\n({cond})")
    _intestazione(ws, 4, intestazioni, [8, 62, 18, 18, 18, 18])
    ws.row_dimensions[4].height = 32

    dv = _validazione_1_7(ws)
    riga = 5
    for chiave, testo in ITEM:
        ws.cell(row=riga, column=1, value=chiave)
        ws.cell(row=riga, column=2, value=testo).alignment = A_CAPO
        for col in range(3, 7):
            cella = ws.cell(row=riga, column=col)
            cella.fill = SFONDO_DA_COMPILARE
            dv.add(cella)
        riga += 1

    riga += 1
    ws.cell(row=riga, column=2, value="condizione (non modificare)").font = Font(color="999999", size=8)
    for col, cond in enumerate(sequenza, start=3):
        c = ws.cell(row=riga, column=col, value=cond)
        c.fill = SFONDO_FISSO
        c.font = Font(color="999999", size=8)
    riga += 1
    ws.cell(row=riga, column=2, value="mercante (non modificare)").font = Font(color="999999", size=8)
    for col, cond in enumerate(sequenza, start=3):
        c = ws.cell(row=riga, column=col, value=MERCANTI[cond])
        c.fill = SFONDO_FISSO
        c.font = Font(color="999999", size=8)

    ws.freeze_panes = "C5"
    return ws


def foglio_finale(wb, codice, build):
    ws = wb.create_sheet("3 Finale")
    ws["A1"] = f"Sessione {codice} — questionario finale (pagine 6-7)"
    ws["A1"].font = TITOLO
    ws["A2"] = ("Parte A: scrivi la CONDIZIONE (V / N / A / D), non il nome e non la posizione. "
                "La corrispondenza e' nel foglio 2.")
    ws["A2"].alignment = A_CAPO

    _intestazione(ws, 4, ["Domanda", "Risposta", "Valori ammessi"], [62, 22, 40])

    condizioni = ["V", "N", "A", "D"]
    domande_a = [
        ("a1_differenze", "1. Ha notato differenze fra i quattro mercanti?", ["si", "no", "non_so"]),
        ("a3_ordine_1", "3. Il piu' 'vero' (1º)", condizioni),
        ("a3_ordine_2", "3. Secondo", condizioni),
        ("a3_ordine_3", "3. Terzo", condizioni),
        ("a3_ordine_4", "3. Il meno 'vero' (4º)", condizioni),
        ("a4_ostile", "4. Il piu' ostile", condizioni),
        ("a4_amichevole", "4. Il piu' amichevole", condizioni),
        ("a5_ricordava", "5. Quale sembrava ricordare qualcosa di lui", condizioni + ["nessuno"]),
        ("a6_coinvolgente", "6. Trattativa piu' coinvolgente", condizioni + ["nessuno"]),
        ("a7_attribuzione", "7. Il prezzo dipendeva da", ["abilita", "rapporto", "caso", "altro"]),
    ]
    riga = 5
    for chiave, testo, ammessi in domande_a:
        ws.cell(row=riga, column=1, value=testo).alignment = A_CAPO
        cella = ws.cell(row=riga, column=2)
        cella.fill = SFONDO_DA_COMPILARE
        _validazione_elenco(ws, ammessi).add(cella)
        ws.cell(row=riga, column=3, value=" / ".join(ammessi)).alignment = A_CAPO
        ws.cell(row=riga, column=4, value=chiave).font = Font(color="999999", size=8)
        riga += 1

    riga += 1
    ws.cell(row=riga, column=1, value="Parte B — la meccanica (scala 1-7)").font = TITOLO
    riga += 1
    dv = _validazione_1_7(ws)
    for chiave, testo in ITEM_B:
        ws.cell(row=riga, column=1, value=testo).alignment = A_CAPO
        cella = ws.cell(row=riga, column=2)
        cella.fill = SFONDO_DA_COMPILARE
        dv.add(cella)
        ws.cell(row=riga, column=4, value=chiave).font = Font(color="999999", size=8)
        riga += 1

    ws.freeze_panes = "A5"
    return ws


def foglio_aperte(wb, codice, build):
    ws = wb.create_sheet("4 Aperte")
    ws["A1"] = f"Sessione {codice} — risposte aperte"
    ws["A1"].font = TITOLO
    ws["A2"] = "Trascrivere alla lettera, senza correggere ne' sintetizzare."

    _intestazione(ws, 4, ["Domanda", "Risposta"], [46, 90])
    voci = []
    for pos, cond in enumerate(SEQUENZE[build], start=1):
        voci.append(f"{pos}. {NOMI[cond]} ({cond}) — che impressione ti ha fatto")
    voci += [
        "Finale 2 — in che cosa erano diversi",
        "Finale 8 — qualcosa che non tornava",
        "Finale B — cosa cambieresti",
        "Finale B — altro",
        "Note del conduttore",
    ]
    for i, testo in enumerate(voci):
        riga = 5 + i
        ws.cell(row=riga, column=1, value=testo).alignment = A_CAPO
        cella = ws.cell(row=riga, column=2)
        cella.fill = SFONDO_DA_COMPILARE
        cella.alignment = A_CAPO
        ws.row_dimensions[riga].height = 46

    ws.freeze_panes = "A5"
    return ws


def crea(codice: str, build: int, destinazione: Path | None = None) -> Path:
    if build not in SEQUENZE:
        raise SystemExit(f"build non valida: {build} (attese 1-4)")

    if destinazione is None:
        destinazione = RADICE / "sessioni" / codice / "dati"
    destinazione.mkdir(parents=True, exist_ok=True)
    percorso = destinazione / f"dati_{codice}.xlsx"
    if percorso.exists():
        raise SystemExit(f"esiste gia': {percorso}  (non lo sovrascrivo)")

    wb = Workbook()
    wb.remove(wb.active)
    foglio_sessione(wb, codice, build)
    foglio_questionari(wb, codice, build)
    foglio_finale(wb, codice, build)
    foglio_aperte(wb, codice, build)
    wb.save(percorso)
    return percorso


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("uso: python scripts/crea_modulo_dati.py <codice> <build 1-4>")
    p = crea(sys.argv[1], int(sys.argv[2]))
    print(f"creato: {p}")
